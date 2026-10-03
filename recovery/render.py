from __future__ import annotations

import io
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .formats import Format
from .util import amount_in_words, fmt_date, inr, parse_date

BLANK = "________"
DEFAULT_BANK = "Union Bank of India"
DEFAULT_HQ = "Union Bank Bhavan, 239 Vidhan Bhavan Marg, Nariman Point, Mumbai 400021"


def first_guarantor(acc: dict) -> str:
    g = (acc.get("guarantors") or "").strip()
    return g.split(";")[0].strip() if g else ""


def base_context(acc: dict, case: dict | None, settings: dict) -> dict:
    out = acc.get("outstanding") or 0
    npa = parse_date(acc.get("as_on"))
    ctx = {
        "bank_name": settings.get("bank_name") or DEFAULT_BANK,
        "bank_hq": settings.get("bank_hq") or DEFAULT_HQ,
        "account_no": acc.get("account_no", ""),
        "borrower": acc.get("borrower", ""),
        "constitution": acc.get("constitution", ""),
        "address": acc.get("address", ""),
        "borrower_address": f"{acc.get('borrower', '')}, {acc.get('address', '')}",
        "facility": acc.get("facility", ""),
        "limit_amt": inr(acc.get("limit_amt") or 0),
        "outstanding": inr(out),
        "outstanding_words": amount_in_words(out),
        "as_on": fmt_date(npa),
        "security_desc": acc.get("security_desc", ""),
        "branch": acc.get("branch", ""),
        "first_guarantor": first_guarantor(acc),
        "officer": (case or {}).get("authorized_officer") or settings.get("officer", ""),
        "designation": (case or {}).get("designation") or settings.get("designation", ""),
        "today": fmt_date(date.today()),
    }
    return ctx


def resolve_defaults(fmt: Format, ctx: dict) -> dict:
    """Initial value for each editable field of the format (strings, or bool for checkboxes)."""
    vals: dict = {}
    for key, _label, default in fmt.fields:
        if isinstance(default, bool):
            vals[key] = default
        elif default == "today":
            vals[key] = ctx["today"]
        elif isinstance(default, str) and default in ctx:
            vals[key] = ctx[default]
        else:
            vals[key] = default or ""
    return vals


def merge(fmt: Format, ctx: dict, values: dict) -> list[str]:
    """Return the template lines after conditionals and merge-field substitution."""
    data = {**ctx, **{k: v for k, v in values.items() if not isinstance(v, bool)}}
    flags = {k for k, v in values.items() if isinstance(v, bool) and v}
    lines = []
    for raw in fmt.body.split("\n"):
        m = re.match(r"^\?(\w+)\|\s?(.*)$", raw)
        if m:
            if m.group(1) not in flags:
                continue
            raw = m.group(2)
        lines.append(re.sub(r"\{\{(\w+)\}\}", lambda mm: str(data.get(mm.group(1)) or BLANK), raw))
    return lines


def to_text(lines: list[str]) -> str:
    out = []
    for ln in lines:
        if ln.startswith("|"):
            out.append("  ".join(c.strip() for c in ln.strip("|").split("|")))
        else:
            out.append(re.sub(r"^(#+|>>|\*\*)\s*", "", ln).replace("**", ""))
    return "\n\n".join(out)


def _bold_runs(par, text: str, bold_all=False):
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for p in parts:
        if not p:
            continue
        if p.startswith("**") and p.endswith("**"):
            par.add_run(p[2:-2]).bold = True
        else:
            r = par.add_run(p)
            r.bold = bold_all


def to_docx(fmt: Format, lines: list[str]) -> bytes:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10.5)
    if fmt.draft:
        p = doc.add_paragraph()
        r = p.add_run("DRAFT - wording not taken from the circular; conform to the SARFAESI Manual form before issue.")
        r.bold = True
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            ncols = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncols)
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(ncols):
                    cell = t.cell(ri, ci)
                    cell.text = row[ci] if ci < len(row) else ""
                    if ri == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        i += 1
        if ln.startswith("##"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _bold_runs(p, ln[2:].strip(), True)
        elif ln.startswith("#"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _bold_runs(p, ln[1:].strip(), True)
            for r in p.runs:
                r.font.size = Pt(12)
        elif ln.startswith(">>"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _bold_runs(p, ln[2:].strip())
        elif ln.startswith("** "):
            p = doc.add_paragraph()
            _bold_runs(p, ln[3:].strip(), True)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _bold_runs(p, ln)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
